from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path(r"C:\Users\Eclipse\Desktop\7 семестр\Програмування мобільних смарт-систем\ЛР-1_МПСС_Гончаров_Вадим_ОІ-43(1).docx")
TARGET = Path(r"C:\Users\Eclipse\Desktop\7 семестр\Програмування мобільних смарт-систем\ЛР-1_МПСС_Гончаров_Вадим_ОІ-43_final.docx")


REPLACEMENTS = {
    "Для виконання серії лабораторних робіт обрано варіант №4": (
        "У межах серії лабораторних робіт за основу було обрано варіант №4, Sleep Tracker Simulator. "
        "Функціональність цього Android-застосунку поступово розширюватиметься у наступних роботах. "
        "Під час виконання першої лабораторної роботи було сформовано базовий сенсорний рівень: отримання "
        "даних акселерометра, відображення координат X, Y і Z, визначення початкового стану рухової "
        "активності та перевірка результату через Virtual Sensors."
    ),
    "Sleep Tracker Simulator у поточній реалізації": (
        "На цьому етапі Sleep Tracker Simulator використовується як навчальний симулятор моніторингу "
        "рухової активності пристрою під час умовного сну. Застосунок не виконує медичного визначення "
        "реальних фаз сну. Реалізований сенсорний рівень є початковою частиною проєкту, на основі якої "
        "можна продовжити роботу в наступних лабораторних."
    ),
    "Після створення проєкту початкову конфігурацію": (
        "Після створення проєкту було підключено модуль application і підтримку Jetpack Compose. Далі у "
        "Gradle додано AndroidX Lifecycle, ViewModel, Material 3 та Kotlin Coroutines, необхідні для "
        "інтерфейсу і передавання сенсорних даних. У процесі реалізації код поступово розподілено між "
        "пакетами core/model, sensor, ui/dashboard та ui/theme. Спочатку в core/model створено моделі "
        "вимірювань. Після цього в sensor розміщено компоненти доступу до акселерометра, а під час "
        "побудови екрана його стан і Compose-компоненти додано до ui/dashboard. Пакет ui/theme залишено "
        "для кольорів та оформлення. MainActivity використано як точку створення залежностей і запуску Compose."
    ),
    "Далі потік станів було перетворено на StateFlow": (
        "Після обробки сенсорних подій потрібно було надати інтерфейсу стан, за яким він оновлюватиме "
        "показники. Тому потік було перетворено на StateFlow у viewModelScope. Режим WhileSubscribed "
        "використано для припинення отримання даних, коли активного екрана немає. На цьому кроці "
        "сенсорний шар було поєднано зі станом інтерфейсу без перенесення викликів SensorManager у "
        "Compose-код. Відповідний фрагмент наведено в лістингу 3."
    ),
    "Самого значення |a| недостатньо": (
        "Для Sleep Tracker Simulator самого значення |a| недостатньо, оскільки воно не дозволяє "
        "відрізнити статичний нахил від руху. Під час повороту нерухомого пристрою проєкції сили тяжіння "
        "на осі X, Y і Z змінюються, однак сумарний модуль після стабілізації залишається близьким до "
        "9,81 м/с². Тому для визначення рухової активності аналізується відстань між двома послідовними "
        "тривимірними векторами:"
    ),
    "Отримана зміна вектора згладжується": (
        "Обчислена зміна вектора може містити короткі коливання, тому перед порівнянням із порогами її "
        "згладжено з коефіцієнтом 0,25. Перехід зі STILL у MOVEMENT виконується після досягнення "
        "0,45 м/с², а повернення у STILL відбувається лише після зниження інтенсивності до 0,20 м/с². "
        "Різні пороги утворюють гістерезис і не дають стану постійно перемикатися через шум біля однієї межі."
    ),
    "У лабораторній роботі створено Android-застосунок": (
        "У результаті виконання лабораторної роботи було створено Android-застосунок Sleep Tracker "
        "Simulator, який отримує показники віртуального акселерометра та відображає їх у Jetpack Compose. "
        "Під час роботи було налаштовано доступ до SensorManager, реєстрацію SensorEventListener і "
        "перетворення отриманих подій у Flow. Окреме джерело даних, класифікатор і SensorViewModel "
        "формувалися послідовно в міру переходу від зчитування координат до їх відображення на екрані."
    ),
    "Для розрізнення нерухомого стану та руху": (
        "Для розрізнення нерухомого стану та руху було проаналізовано зміну тривимірного вектора "
        "прискорення, після чого додано згладжування і гістерезис. Під час перевірки через Virtual Sensors "
        "підтверджено відповідність координат X, Y і Z між емулятором та застосунком. Статичний нахил "
        "після стабілізації залишав систему у STILL, швидка зміна Device Pose переводила її у MOVEMENT, "
        "а після припинення руху відбувалося повернення у STILL."
    ),
    "callbackFlow та awaitClose відключають": (
        "Окрему увагу було приділено керуванню життєвим циклом сенсора. callbackFlow, awaitClose і "
        "WhileSubscribed припиняють зчитування, коли інтерфейс більше не споживає дані. Це важливо, "
        "оскільки слухач акселерометра не повинен залишатися активним після завершення роботи з екраном. "
        "Отриманий результат виконує завдання першої лабораторної роботи та формує основу для подальшого "
        "розвитку Sleep Tracker Simulator без медичних тверджень про визначення фаз сну."
    ),
}


def set_run_font(run, name="Times New Roman", size=14):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)


def replace_paragraph(paragraph, text):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    set_run_font(run)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    shutil.copy2(SOURCE, TARGET)
    doc = Document(TARGET)

    found = set()
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        for prefix, replacement in REPLACEMENTS.items():
            if text.startswith(prefix):
                replace_paragraph(paragraph, replacement)
                found.add(prefix)
                break

        if text.startswith("|a| ="):
            replace_paragraph(paragraph, "[[WORD_EQUATION_1]]")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif text.startswith("Δa ="):
            replace_paragraph(paragraph, "[[WORD_EQUATION_2]]")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    missing = set(REPLACEMENTS) - found
    if missing:
        raise RuntimeError(f"Paragraphs not found: {sorted(missing)}")

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True

    doc.save(TARGET)


if __name__ == "__main__":
    main()
