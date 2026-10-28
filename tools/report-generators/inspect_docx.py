from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(path)
    data = {
        "path": str(path),
        "paragraphs": [
            {"index": i, "style": p.style.name if p.style else None, "text": p.text}
            for i, p in enumerate(doc.paragraphs)
            if p.text.strip()
        ],
        "tables": [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in doc.tables
        ],
        "sections": [
            {
                "width": s.page_width,
                "height": s.page_height,
                "top": s.top_margin,
                "bottom": s.bottom_margin,
                "left": s.left_margin,
                "right": s.right_margin,
            }
            for s in doc.sections
        ],
        "inline_shapes": len(doc.inline_shapes),
        "styles": [s.name for s in doc.styles if s.type == 1],
    }
    print(json.dumps(data, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
