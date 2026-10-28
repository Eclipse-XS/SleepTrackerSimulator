from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
REF = ROOT / "deliverables" / "ЛР-5_ПМСС_Гончаров_Вадим_ОІ-43.docx"
OUT = ROOT / "deliverables"

def font(run, name="Times New Roman", size=14, bold=False):
    run.font.name = name; run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr; el = OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement(f"w:{edge}"); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4"); e.set(qn("w:color"),"D9D9D9"); el.append(e)
    tblPr.append(el)

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,"1F4E78"); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: font(r,size=10,bold=True); r.font.color.rgb=RGBColor(255,255,255)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx%2: shade(cells[i],"EAF2F8")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: font(r,size=9)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph()

def reset_reference():
    doc=Document(REF)
    body=doc._element.body; sect=body.sectPr
    for child in list(body):
        if child is not sect: body.remove(child)
    styles=doc.styles
    for style_name,size in (("Normal",14),("Title",18),("Heading 1",16),("Heading 2",14)):
        s=styles[style_name]; s.font.name="Times New Roman"; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    return doc

def title(doc, number, topic):
    for text,bold,size,space in [
        ("Міністерство освіти і науки України",False,14,0),
        ("ЛАБОРАТОРНА РОБОТА",True,18,0),
        (f"№ {number}",True,18,0),
        ("з дисципліни «Програмування мобільних смарт-систем»",False,14,0),
        (topic,True,14,0),
    ]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(space)
        font(p.add_run(text),size=size,bold=bold)
    doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(p.add_run("Виконав: студент групи ОІ-43\nГончаров Вадим"),size=14)
    for _ in range(5): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("2026"),size=14)
    doc.add_page_break()

def h(doc,text,level=1): doc.add_heading(text,level=level)
def para(doc,text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.first_line_indent=Cm(1.25); p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(6)
    font(p.add_run(text)); return p
def code(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.6); p.paragraph_format.space_after=Pt(6)
    for line in text.splitlines(): font(p.add_run(line+"\n"),name="Consolas",size=9)

def lab6():
    d=reset_reference(); title(d,6,"Інтеграція моделей машинного навчання у мобільний додаток")
    h(d,"1 Мета та індивідуальне завдання"); para(d,"Мета роботи — реалізувати локальну багатокласову класифікацію рухової активності за вікном акселерометра й інтегрувати результат із dashboard та актуатором. Індивідуальний домен — освітній симулятор спокою і сну. Система не є медичним засобом.")
    h(d,"2 Теоретичні основи Edge AI"); para(d,"On-device inference виконується без передавання сирого вікна до хмари. Переваги: низька затримка, офлайн-робота та менше розкриття сенсорних даних. Обмеження: ресурси телефона, необхідність контролю життєвого циклу та складність репрезентативної валідації.")
    h(d,"3 Дані та відтворюваність"); para(d,"Скрипт ml/train_activity_model.py з seed 4319 генерує синтетичні тривісні сигнали. REST містить малий шум, RESTLESS — помірні періодичні рухи, ACTIVE — сильніші коливання та імпульси. Навчальна і тестова вибірки генеруються різними seed і не перетинаються. Синтетичний test set не замінює польову валідацію.")
    table(d,["Параметр","Значення"],[("Train","3600 вікон"),("Test","900 вікон"),("Window","20 × 3 float32"),("Класи","REST, RESTLESS, ACTIVE"),("Seed","4319"),("TensorFlow","2.20.0")],[5,11])
    h(d,"4 Попередня обробка і модель"); para(d,"Вхід має форму [1,20,3]. Усередині графа від кожної осі віднімається середнє за часом. Далі застосовуються Flatten, Dense(32, ReLU) і Dense(3, Softmax). Вихід [1,3] є розподілом імовірностей.")
    code(d,"Input [1,20,3] float32\nMean centering over 20 samples\nFlatten -> Dense(32, ReLU) -> Dense(3, Softmax)\nOutput [1,3] float32")
    h(d,"5 Результати навчання"); para(d,"Фактичний запуск 19 жовтня 2026 року дав test loss 0,177139 і accuracy 0,972222. Файл activity_classifier.tflite має 11368 байт.")
    table(d,["Фактичний клас","REST","RESTLESS","ACTIVE"],[("REST",300,0,0),("RESTLESS",18,282,0),("ACTIVE",0,7,293)],[4.5,3.5,3.5,3.5])
    h(d,"6 Android інтеграція"); para(d,"ActivityClassifier завантажує asset один раз, перевіряє форми тензорів і відповідність labels. SensorWindowBuffer повертає хронологічне rolling window. MlViewModel запускає inference через Dispatchers.Default і закриває Interpreter у onCleared.")
    code(d,"val window = buffer.add(sample) ?: return\nval prediction = withContext(Dispatchers.Default) {\n    classifier.classify(window)\n}")
    h(d,"7 AI Insights і актуатор"); para(d,"Картка AI Insights показує клас, confidence і виміряну тривалість inference. Перехід у ACTIVE з confidence не менше 0,75 створює одиничний vibration pulse, якщо автоматичний режим увімкнений і застосунок у foreground. Подія записується до Room.")
    h(d,"8 Тестування"); table(d,["ID","Перевірка","Результат"],[("ML-01","Rolling window і clear","PASS, unit"),("ML-02","Реальний TFLite REST inference","PASS, instrumentation"),("ML-03","3 скінченні probabilities","PASS, instrumentation"),("ML-04","UI на різних Virtual Sensors","Потребує ручних screenshots")],[2,8,6])
    h(d,"9 Скріншоти та докази"); para(d,"Структура проєкту містить app/src/main/assets/activity_classifier.tflite і activity_labels.txt. Автоматичні тести не є заміною screenshots зміни прогнозу в Extended Controls. Такі screenshots не позначаються виконаними без фактичного capture.")
    h(d,"10 Висновки"); para(d,"Реалізовано справжній локальний TFLite inference з фіксованим контрактом тензорів, кільцевим буфером і керованим життєвим циклом. Основне обмеження — синтетичні дані; для практичного використання потрібен реальний датасет і валідація на різних пристроях.")
    path=OUT/"ЛР-6_ПМСС_Гончаров_Вадим_ОІ-43.docx"; d.save(path); return path

def lab7():
    d=reset_reference(); title(d,7,"Мініпроєкт цілісної смарт системи SleepTrackerSimulator")
    h(d,"1 Мета і концепція"); para(d,"Мета — об’єднати SensorManager, Edge AI, Room, Firebase, WorkManager, актуатори та Mock BLE в один перевірюваний pipeline. Система зберігає роботу без мережі та явно показує pending або error state.")
    h(d,"2 Функціональні вимоги"); table(d,["Компонент","Функція","Стан"],[("Sensor","XYZ і movement state","Реалізовано"),("BLE","Scan Connect Read Write Notify","Mock transport"),("ML","REST RESTLESS ACTIVE","Реальний TFLite"),("Room","History events outbox","Schema v3"),("Cloud","Owner scoped Firebase sync","Durable retry"),("Actuator","Vibration torch manual remote auto","Реалізовано")],[4,8,4])
    h(d,"3 Архітектура і життєвий цикл"); para(d,"MainActivity виконує manual DI. Sensor listener існує лише у foreground. ViewModel-и володіють станом модулів. Room singleton є application-scoped, ML Interpreter — MlViewModel-scoped, BLE connector — BleViewModel-scoped, WorkManager worker — worker-scoped. Foreground і worker drainers серіалізовані mutex-ом.")
    code(d,"SensorManager -> SensorViewModel -> CircularBuffer -> TFLite\n                 |-> Room history\n                 |-> Room outbox -> WorkManager -> Firebase\nTFLite ACTIVE -> ActuatorViewModel -> vibration -> Room event\nVirtualSleepBand -> BleViewModel -> rest estimation and cloud payload")
    h(d,"4 Потік даних"); para(d,"Callback акселерометра формує SensorSample з monotonic timestamp. MovementClassifier обчислює intensity, SensorViewModel публікує StateFlow. LocalHistoryRecorderViewModel робить періодичний запис. MlViewModel накопичує 20 samples. CloudSyncViewModel раз на 5 секунд створює stable operation і додає його до Room outbox до мережевої доставки.")
    h(d,"5 Відмовостійкість"); para(d,"За відсутності мережі FIFO-операції залишаються в Room. Межа telemetry queue — 10000. Після заповнення нова телеметрія отримує явний backpressure без eviction. Один унікальний END_SESSION дозволений як окрема terminal operation навіть для повної telemetry queue; повторний terminal запис відхиляється. Після відновлення мережі WorkManager продовжує доставку.")
    h(d,"6 Безпека"); para(d,"Firebase rules обмежують дані owner UID, перевіряють типи та забороняють невідомі поля. Worker повторно перевіряє UID. Backup і device transfer вимкнені. Mock BLE не потребує зайвих Bluetooth або location permissions. Це не усуває потреби у production privacy flow для export і delete.")
    h(d,"7 Dashboard"); para(d,"Dashboard показує акселерометр XYZ, movement summary, AI Insights, BLE connection, heart rate, rest estimate, cloud pending state та quick actuator controls. Analytics читає Room Flow. Відсутні актуальні ручні screenshots не підміняються штучними зображеннями.")
    h(d,"8 Тест кейси"); table(d,["ID","Вхід","Очікування","Факт","Статус"],[("E2E-01","20 REST samples","REST prediction","Real interpreter test","PASS"),("E2E-02","Offline uploads","FIFO retained","Unit and Room tests","PASS"),("E2E-03","Full outbox then END","Terminal retained","Instrumentation test","PASS"),("E2E-04","Reconnect","FIFO drain","Component integration tests","PASS"),("E2E-05","Virtual Sensors UI","Visible transitions","No screenshot evidence","BLOCKED"),("E2E-06","Profiler soak","No retained listeners","Not executed in Profiler","BLOCKED"),("E2E-07","Full video","Complete scenario","Not recorded","BLOCKED")],[1.7,4,4.2,4.2,2])
    h(d,"9 Продуктивність і ресурси"); para(d,"UI показує реальну тривалість кожного inference. Модель займає 11368 байт без урахування runtime. Частота inference обмежена надходженням StateFlow samples; Interpreter не створюється під час recomposition. Повний P95, heap-growth та battery soak потребують окремого вимірювального запуску й не заявляються виконаними.")
    h(d,"10 Обмеження і масштабування"); para(d,"Ключові обмеження: синтетична ML-вибірка, Mock BLE замість фізичного transport, відсутність довгого profiler soak, відео та post-change screenshot-пакета, вимкнена release optimization. Масштабування потребує реального датасету, versioned model rollout, physical BLE implementation, data export/delete і release hardening.")
    h(d,"11 Висновки"); para(d,"Цілісний навчальний pipeline реалізовано і підтримано автоматичними тестами компонентів та реального TFLite interpreter-а. ЛР7 не можна назвати повністю завершеною щодо доказів захисту, поки не виконано ручну демонстрацію, profiler session і відеозапис.")
    path=OUT/"ЛР-7_ПМСС_Гончаров_Вадим_ОІ-43.docx"; d.save(path); return path

if __name__ == "__main__":
    print(lab6()); print(lab7())
