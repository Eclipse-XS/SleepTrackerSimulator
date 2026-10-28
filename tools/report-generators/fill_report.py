from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SOURCE = Path(r"C:\Users\Eclipse\Desktop\7 семестр\Програмування мобільних смарт-систем\ЛР-1_МПСС_Гончаров_Вадим_ОІ-43.docx")
TARGET = Path(r"C:\Users\Eclipse\Desktop\7 семестр\Програмування мобільних смарт-систем\ЛР-1_МПСС_Гончаров_Вадим_ОІ-43_final.docx")
TMP = TARGET.with_name(TARGET.stem + ".tmp.docx")
IMAGES = {
    "studio": Path(r"C:\Users\Eclipse\AppData\Local\Temp\codex-clipboard-bb53a396-c48e-4f0e-bc76-06a0e74d6dce.png"),
    "project": Path(r"C:\Users\Eclipse\AppData\Local\Temp\codex-clipboard-9fe18f3b-77f3-4792-a4d7-e966d75f12e8.png"),
    "still": Path(r"C:\Users\Eclipse\AppData\Local\Temp\codex-clipboard-725b47a0-a373-42b5-b057-7c105bb1f433.png"),
    "tilt": Path(r"C:\Users\Eclipse\AppData\Local\Temp\codex-clipboard-d895db4c-7982-4237-9958-77a53c2c1caf.png"),
    "movement": Path(r"C:\Users\Eclipse\AppData\Local\Temp\codex-clipboard-3e8a2758-c9c1-4f03-8500-820ad56ec56e.png"),
}


def set_font(run, name="Times New Roman", size=14, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def body(doc, text, first_line=True, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(after)
    if first_line:
        pf.first_line_indent = Cm(1.25)
    r = p.add_run(text)
    set_font(r)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 9)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=14, bold=True)
    return p


def label_para(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    r = p.add_run(label)
    set_font(r, bold=True)
    r = p.add_run(text)
    set_font(r)
    return p


def list_item(doc, text, number=None, level=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(3)
    pf.left_indent = Cm(0.75 + level * 0.6)
    pf.first_line_indent = Cm(-0.5)
    marker = f"{number}. " if number is not None else "• "
    r = p.add_run(marker + text)
    set_font(r)
    return p


def figure(doc, image_path, caption, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(6)
    c.paragraph_format.keep_together = True
    r = c.add_run(caption)
    set_font(r, size=12, italic=True)


def code_listing(doc, title, code):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, size=12, bold=True)

    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16.2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F3F3")
    tc_pr.append(shd)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        borders.append(el)
    tc_pr.append(borders)
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", "100"), ("left", "140"), ("bottom", "100"), ("right", "140")):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), value)
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tc_pr.append(margins)
    cp = cell.paragraphs[0]
    cp.paragraph_format.line_spacing = 1.0
    cp.paragraph_format.space_after = Pt(0)
    run = cp.add_run(code.strip())
    set_font(run, name="Consolas", size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def formula(doc, parts):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    source = "".join(text for text, _ in parts)
    marker = "[[WORD_EQUATION_1]]" if source.startswith("|a|") else "[[WORD_EQUATION_2]]"
    r = p.add_run(marker)
    set_font(r, size=14, italic=True)


def remove_after_title(doc):
    marker = next(p for p in doc.paragraphs if p.text.strip() == "Львів-2026")
    body_element = doc._element.body
    children = list(body_element)
    marker_index = children.index(marker._element)
    for element in children[marker_index + 1:]:
        if element.tag != qn("w:sectPr"):
            body_element.remove(element)


def main():
    for path in IMAGES.values():
        if not path.exists():
            raise FileNotFoundError(path)

    doc = Document(SOURCE)
    remove_after_title(doc)

    # Academic body formatting follows the supplied reports; the original title page is untouched.
    heading(doc, "Мета роботи")
    body(doc, "Ознайомитися з середовищем Android Studio, принципами побудови інтерфейсу на Jetpack Compose та механізмом отримання даних від SensorManager в ОС Android.")

    heading(doc, "Концепція наскрізного завдання")
    body(doc, "У межах серії лабораторних робіт за основу було обрано варіант №4, Sleep Tracker Simulator. Функціональність цього Android-застосунку поступово розширюватиметься у наступних роботах. Під час виконання першої лабораторної роботи було сформовано базовий сенсорний рівень: отримання даних акселерометра, відображення координат X, Y і Z, визначення початкового стану рухової активності та перевірка результату через Virtual Sensors.")
    body(doc, "На цьому етапі Sleep Tracker Simulator використовується як навчальний симулятор моніторингу рухової активності пристрою під час умовного сну. Застосунок не виконує медичного визначення реальних фаз сну. Реалізований сенсорний рівень є початковою частиною проєкту, на основі якої можна продовжити роботу в наступних лабораторних.")

    heading(doc, "Завдання на лабораторну роботу")
    list_item(doc, "Створити Android-проєкт для реалізації Sleep Tracker Simulator.", 1)
    list_item(doc, "Реалізувати графічний інтерфейс для відображення значень X, Y і Z акселерометра, модуля прискорення, графічного індикатора положення та руху, а також поточного стану STILL або MOVEMENT.", 2)
    list_item(doc, "Налаштувати роботу із SensorManager: отримати Sensor.TYPE_ACCELEROMETER, зареєструвати SensorEventListener, реалізувати обробку onSensorChanged і припиняти зчитування після завершення споживання даних.", 3)
    list_item(doc, "Реалізувати порогову логіку переходу між станами STILL і MOVEMENT.", 4)
    list_item(doc, "Протестувати застосунок через Extended Controls, Virtual Sensors і Device Pose.", 5)
    list_item(doc, "Перевірити реакцію системи на нерухомий стан, статичний нахил і швидку зміну положення пристрою.", 6)

    heading(doc, "Технічний стек")
    list_item(doc, "Мова програмування: Kotlin.")
    list_item(doc, "Мінімальна версія Android: Android 12, API 31 і вище.")
    list_item(doc, "Інтерфейс користувача: Jetpack Compose.")
    list_item(doc, "Середовище розробки: Android Studio.")
    list_item(doc, "Сенсорний API: SensorManager, SensorEventListener і Sensor.TYPE_ACCELEROMETER.")
    list_item(doc, "Робота з потоками даних: Kotlin Coroutines, Flow і StateFlow.")
    list_item(doc, "Архітектурний компонент: ViewModel.")
    list_item(doc, "Середовище тестування: Android Emulator, Pixel 10 Pro, Android 16, API 36.1.")
    list_item(doc, "Імітація сенсорів: Extended Controls, Virtual Sensors.")

    heading(doc, "Індивідуальне завдання")
    body(doc, "Варіант №4, Sleep Tracker Simulator: моніторинг фаз сну на основі мікрорухів пристрою.")
    body(doc, "У межах першої лабораторної роботи вихідне формулювання трактується як навчальний аналіз рухової активності пристрою. Застосунок не визначає медичні фази сну, а розрізняє два базові стани: STILL і MOVEMENT. Для цього використовуються поточні показники акселерометра та зміна тривимірного вектора між послідовними вимірюваннями.")

    heading(doc, "Обґрунтування порогових значень")
    body(doc, "Модуль вектора прискорення використовується для загальної оцінки вимірювання акселерометра:")
    formula(doc, [("|a| = √(x", False), ("2", True), (" + y", False), ("2", True), (" + z", False), ("2", True), (")", False)])
    body(doc, "Для Sleep Tracker Simulator самого значення |a| недостатньо, оскільки воно не дозволяє відрізнити статичний нахил від руху. Під час повороту нерухомого пристрою проєкції сили тяжіння на осі X, Y і Z змінюються, однак сумарний модуль після стабілізації залишається близьким до 9,81 м/с². Тому для визначення рухової активності аналізується відстань між двома послідовними тривимірними векторами:")
    formula(doc, [("Δa = √((xₜ − xₜ₋₁)", False), ("2", True), (" + (yₜ − yₜ₋₁)", False), ("2", True), (" + (zₜ − zₜ₋₁)", False), ("2", True), (")", False)])
    body(doc, "Обчислена зміна вектора може містити короткі коливання, тому перед порівнянням із порогами її згладжено з коефіцієнтом 0,25. Перехід зі STILL у MOVEMENT виконується після досягнення 0,45 м/с², а повернення у STILL відбувається лише після зниження інтенсивності до 0,20 м/с². Різні пороги утворюють гістерезис і не дають стану постійно перемикатися через шум біля однієї межі.")

    heading(doc, "Хід роботи")
    body(doc, "Виконання роботи було поділено на шість послідовних етапів. Спочатку було підготовлено Android-проєкт і підключено необхідні залежності. Після цього реалізовано отримання показників акселерометра, додано визначення станів STILL і MOVEMENT та передавання результату до інтерфейсу. На завершальних етапах було створено віртуальний пристрій, налаштовано Virtual Sensors і перевірено реакцію застосунку на нерухоме положення, нахил та швидку зміну Device Pose.")

    heading(doc, "Етап 1 Створення та початкове налаштування Android проєкту", 2)
    body(doc, "На початковому етапі необхідно було підготувати середовище, у якому надалі створювався і тестувався мобільний застосунок. Для цього використано Android Studio, оскільки воно поєднує засоби створення проєкту, налаштування Gradle, написання Kotlin-коду та запуск Android Emulator. Офіційну сторінку, з якої доступне завантаження середовища, показано на рис. 1.")
    figure(doc, IMAGES["studio"], "Рис. 1. Офіційна сторінка завантаження Android Studio")
    body(doc, "Після підготовки середовища було створено новий проєкт SleepTrackerSimulator. Як початковий шаблон обрано No Activity, оскільки інтерфейс і точку входу потрібно було сформувати відповідно до завдання лабораторної роботи. Мовою розробки встановлено Kotlin, а мінімальною версією системи обрано API 31. Вибрані параметри створення проєкту наведено на рис. 2.")
    figure(doc, IMAGES["project"], "Рис. 2. Створення проєкту SleepTrackerSimulator з мінімальним API 31")
    body(doc, "Після створення проєкту було підключено модуль application і підтримку Jetpack Compose. Далі у Gradle додано AndroidX Lifecycle, ViewModel, Material 3 та Kotlin Coroutines, необхідні для інтерфейсу і передавання сенсорних даних. У процесі реалізації код поступово розподілено між пакетами core/model, sensor, ui/dashboard та ui/theme. Спочатку в core/model створено моделі вимірювань. Після цього в sensor розміщено компоненти доступу до акселерометра, а під час побудови екрана його стан і Compose-компоненти додано до ui/dashboard. Пакет ui/theme залишено для кольорів та оформлення. MainActivity використано як точку створення залежностей і запуску Compose.")

    heading(doc, "Етап 2 Реалізація отримання даних акселерометра", 2)
    body(doc, "Після підготовки базової структури проєкту наступним кроком стало підключення акселерометра. Для отримання його показників в Android використовується системний сервіс SensorManager. Щоб робота із системним API не виконувалася безпосередньо на екрані, спочатку було створено контракт AccelerometerDataSource. Він визначає операцію observe, через яку інші компоненти отримують потік подій акселерометра.")
    body(doc, "Далі було реалізовано клас AndroidAccelerometerDataSource. Під час його створення з контексту отримується SensorManager, після чого запитується стандартний Sensor.TYPE_ACCELEROMETER. Окремо передбачено випадок, коли потрібний сенсор недоступний: замість спроби продовжити роботу формується подія Unavailable. Аналогічно помилка під час реєстрації слухача передається як Error, щоб інтерфейс міг показати відповідний стан.")
    body(doc, "Для передавання повторюваних викликів SensorEventListener у Kotlin Flow використано callbackFlow. Усередині нього було створено слухач, а в методі onSensorChanged додано перевірку типу сенсора та кількості отриманих координат. Після успішної перевірки значення event.values[0], event.values[1] і event.values[2] записуються у SensorSample разом із часовою міткою event.timestamp. Фрагмент реалізації наведено в лістингу 1.")
    code_listing(doc, "Лістинг 1. Отримання вимірювань акселерометра", r'''
override fun observe(): Flow<AccelerometerEvent> = callbackFlow {
    val listener = object : SensorEventListener {
        override fun onSensorChanged(event: SensorEvent) {
            if (event.sensor.type != Sensor.TYPE_ACCELEROMETER ||
                event.values.size < 3
            ) return

            trySend(
                AccelerometerEvent.Measurement(
                    SensorSample(
                        x = event.values[0],
                        y = event.values[1],
                        z = event.values[2],
                        timestampNanos = event.timestamp,
                    ),
                ),
            )
        }

        override fun onAccuracyChanged(sensor: Sensor?, accuracy: Int) = Unit
    }

    manager.registerListener(listener, sensor, SensorManager.SENSOR_DELAY_UI)
    awaitClose { manager.unregisterListener(listener) }
}''')
    body(doc, "Після створення слухача його було зареєстровано з частотою SENSOR_DELAY_UI, якої достатньо для оновлення показників на екрані. Важливо було також припиняти зчитування після завершення використання потоку. Тому в awaitClose викликається unregisterListener. У результаті слухач існує лише тоді, коли застосунок фактично отримує дані, і не продовжує роботу після закриття підписки.")

    heading(doc, "Етап 3 Реалізація визначення рухової активності", 2)
    body(doc, "Після отримання стабільного потоку координат потрібно було перейти від відображення сирих значень до визначення рухової активності. Для цього створено MovementClassifier. Перший отриманий SensorSample використовується як початкова точка, тому для нього інтенсивність дорівнює нулю. Після надходження наступного вимірювання обчислюється різниця за кожною з трьох осей, а потім визначається відстань між поточним і попереднім векторами.")
    body(doc, "Початкове значення зміни може містити короткі стрибки, тому перед порівнянням із порогами його було згладжено. Для цього попередня інтенсивність поступово наближається до нового значення з коефіцієнтом 0,25. Після додавання згладжування було реалізовано окремі умови входу у MOVEMENT і повернення у STILL. Відповідну частину класифікатора наведено в лістингу 2.")
    code_listing(doc, "Лістинг 2. Класифікація станів STILL і MOVEMENT", r'''
val rawIntensity = if (previous == null) {
    0f
} else {
    val deltaX = sample.x - previous.x
    val deltaY = sample.y - previous.y
    val deltaZ = sample.z - previous.z
    sqrt(deltaX * deltaX + deltaY * deltaY + deltaZ * deltaZ)
}
previousSample = sample
smoothedIntensity += smoothingFactor * (rawIntensity - smoothedIntensity)

state = when (state) {
    MovementState.STILL ->
        if (smoothedIntensity >= movementThreshold)
            MovementState.MOVEMENT else MovementState.STILL
    MovementState.MOVEMENT ->
        if (smoothedIntensity <= stillThreshold)
            MovementState.STILL else MovementState.MOVEMENT
}''')
    body(doc, "Після цього було перевірено логіку переходів. У стані STILL класифікатор очікує, доки згладжена зміна не досягне 0,45 м/с². Після переходу в MOVEMENT він не повертається назад відразу, а очікує зменшення показника до 0,20 м/с². Завдяки цьому короткі коливання біля одного порога не спричиняють постійної зміни напису на екрані. Коли нахил пристрою вже не змінюється, різниця між сусідніми векторами зменшується і стан повертається у STILL.")

    heading(doc, "Етап 4 Формування стану та користувацького інтерфейсу", 2)
    body(doc, "Після завершення класифікатора потрібно було передати його результат до Compose-інтерфейсу. Для цього створено SensorViewModel, яка підписується на AccelerometerDataSource. Кожна подія Measurement надходить до MovementClassifier, а отриманий MovementAnalysis записується у стан Active. До появи першого вимірювання використовується Initializing, а для відсутнього сенсора та помилок передбачено Unavailable і Error.")
    body(doc, "Після обробки сенсорних подій потрібно було надати інтерфейсу стан, за яким він оновлюватиме показники. Тому потік було перетворено на StateFlow у viewModelScope. Режим WhileSubscribed використано для припинення отримання даних, коли активного екрана немає. На цьому кроці сенсорний шар було поєднано зі станом інтерфейсу без перенесення викликів SensorManager у Compose-код. Відповідний фрагмент наведено в лістингу 3.")
    code_listing(doc, "Лістинг 3. Перетворення сенсорного потоку на стан інтерфейсу", r'''
val uiState = dataSource.observe()
    .map { event ->
        when (event) {
            is AccelerometerEvent.Measurement ->
                SensorUiState.Active(classifier.analyze(event.sample))
            AccelerometerEvent.Unavailable -> SensorUiState.Unavailable
            is AccelerometerEvent.Error -> SensorUiState.Error(event.message)
        }
    }
    .onStart {
        classifier.reset()
        emit(SensorUiState.Initializing)
    }
    .stateIn(
        scope = viewModelScope,
        started = SharingStarted.WhileSubscribed(stopTimeoutMillis = 0),
        initialValue = SensorUiState.Initializing,
    )''')
    body(doc, "Коли передавання стану було налаштовано, наступним кроком стала побудова екрана SensorScreen. У верхній частині додано назву Sleep Tracker Simulator і уточнення, що система виконує навчальний моніторинг руху, а не медичне визначення фаз сну. Нижче розміщено картку поточного стану, окремі значення X, Y і Z, модуль прискорення та графічний індикатор положення.")
    body(doc, "Для індикатора координати X і Y було нормалізовано відносно 9,81 м/с² і використано для зміщення точки всередині кола. Навколишня дуга показує відносну інтенсивність руху, а під нею виводиться точне згладжене значення Filtered motion. Після цього було додано різне кольорове оформлення: зелений стан STILL і помаранчевий MOVEMENT. Вертикальне прокручування екрана залишає всі елементи доступними після зміни орієнтації емулятора.")

    heading(doc, "Етап 5 Налаштування Android Emulator та Virtual Sensors", 2)
    body(doc, "Після завершення основної реалізації було підготовлено середовище для її практичної перевірки. Застосунок запущено на віртуальному пристрої Pixel 10 Pro з Android 16 та API 36.1. Далі у вікні емулятора відкрито Extended Controls, розділ Virtual Sensors і вкладку Device Pose. За допомогою X-Rot, Y-Rot і Z-Rot у цьому вікні можна змінювати положення віртуального телефона та одночасно спостерігати показники Accelerometer, Gyroscope і Magnetometer. Початковий нерухомий стан наведено на рис. 3.")
    figure(doc, IMAGES["still"], "Рис. 3. Інтерфейс Sleep Tracker Simulator і Virtual Sensors у нерухомому стані")
    body(doc, "На рис. 3 одночасно видно панель Virtual Sensors і запущений Sleep Tracker Simulator. У панелі сенсорів відображаються X = 0,00, Y = 9,78 і Z = 0,81 м/с². Ті самі округлені значення показані у картках застосунку, а обчислений модуль становить 9,81 м/с². Отже, перше порівняння підтвердило проходження даних від віртуального акселерометра до створеного інтерфейсу без зміни координат.")

    heading(doc, "Етап 6 Тестування роботи застосунку", 2)
    body(doc, "Після перевірки початкового стану було змінено положення пристрою повзунками Device Pose. Спочатку виконано повільний нахил, після якого телефон залишено нерухомим. У цьому положенні Virtual Sensors показує X = 3,47, Y = 5,18 і Z = 7,57 м/с². Відповідний результат роботи застосунку наведено на рис. 4.")
    figure(doc, IMAGES["tilt"], "Рис. 4. Відповідність показників акселерометра при статичному нахилі пристрою")
    body(doc, "На рис. 4 значення трьох осей у Virtual Sensors збігаються з картками X, Y і Z у застосунку. Водночас після завершення нахилу індикатор повернувся у STILL. Цей сценарій був потрібний для перевірки того, що змінена орієнтація сама по собі не розглядається як безперервний рух: після стабілізації сусідні вектори знову стають близькими.")
    body(doc, "Наступним сценарієм стала швидка зміна Device Pose. Під час різкого переміщення значення координат змінювалися за короткий проміжок часу, згладжена інтенсивність перевищила 0,45 м/с², а картка стану змінилася зі STILL на MOVEMENT. Одночасно зелений колір було замінено на помаранчевий. Отриманий стан показано на рис. 5.")
    figure(doc, IMAGES["movement"], "Рис. 5. Перехід Sleep Tracker Simulator у стан MOVEMENT", width_cm=11.5)
    body(doc, "Рис. 5 також демонструє, що екран залишається придатним до перегляду в альбомній орієнтації. Після припинення зміни положення згладжена інтенсивність почала зменшуватися. Коли вона стала нижчою за 0,20 м/с², система повернулася у STILL. Таким чином, послідовно було перевірено обидва переходи класифікатора та відмінність між короткочасним рухом і вже стабільним нахилом.")
    body(doc, "На завершення виконано завдання Gradle assembleDebug, test і lint. Усі вони завершилися зі статусом BUILD SUCCESSFUL. Запущено шість модульних тестів без помилок, падінь і пропусків. Ними перевірено розрахунок модуля за трьома осями, збереження STILL для незмінного вектора сили тяжіння, перехід у MOVEMENT після перевищення верхнього порога, роботу гістерезису та повернення класифікатора у початковий стан після reset.")

    heading(doc, "Висновок")
    body(doc, "У результаті виконання лабораторної роботи було створено Android-застосунок Sleep Tracker Simulator, який отримує показники віртуального акселерометра та відображає їх у Jetpack Compose. Під час роботи було налаштовано доступ до SensorManager, реєстрацію SensorEventListener і перетворення отриманих подій у Flow. Окреме джерело даних, класифікатор і SensorViewModel формувалися послідовно в міру переходу від зчитування координат до їх відображення на екрані.")
    body(doc, "Для розрізнення нерухомого стану та руху було проаналізовано зміну тривимірного вектора прискорення, після чого додано згладжування і гістерезис. Під час перевірки через Virtual Sensors підтверджено відповідність координат X, Y і Z між емулятором та застосунком. Статичний нахил після стабілізації залишав систему у STILL, швидка зміна Device Pose переводила її у MOVEMENT, а після припинення руху відбувалося повернення у STILL.")
    body(doc, "Окрему увагу було приділено керуванню життєвим циклом сенсора. callbackFlow, awaitClose і WhileSubscribed припиняють зчитування, коли інтерфейс більше не споживає дані. Це важливо, оскільки слухач акселерометра не повинен залишатися активним після завершення роботи з екраном. Отриманий результат виконує завдання першої лабораторної роботи та формує основу для подальшого розвитку Sleep Tracker Simulator без медичних тверджень про визначення фаз сну.")

    # Normalize body defaults without altering the title page's direct formatting.
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(14)

    doc.save(TMP)
    os.replace(TMP, TARGET)


if __name__ == "__main__":
    main()
