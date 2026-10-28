from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Users\Eclipse\AndroidStudioProjects\SleepTrackerSimulator")
OUT_DIR = ROOT / "deliverables"
EVIDENCE = OUT_DIR / "evidence"
WORK = ROOT / ".codex-doc"
OUT = OUT_DIR / "ЛР-5_ПМСС_Гончаров_Вадим_ОІ-43_WORKING.docx"
ER = EVIDENCE / "lab5-er.png"
MVVM = EVIDENCE / "lab5-mvvm.png"
SCREEN = EVIDENCE / "lab5-analytics-final.png"
LOGO = WORK / "lpnu-logo.png"

OUT_DIR.mkdir(parents=True, exist_ok=True)
EVIDENCE.mkdir(parents=True, exist_ok=True)

def font(size=28, bold=False):
    path = Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf")
    return ImageFont.truetype(str(path), size)

def arrow(draw, start, end, width=4):
    draw.line((start, end), fill="#24364B", width=width)
    x1, y1 = start; x2, y2 = end
    if abs(x2-x1) > abs(y2-y1):
        s = 1 if x2 > x1 else -1
        draw.polygon([(x2,y2),(x2-15*s,y2-9),(x2-15*s,y2+9)], fill="#24364B")
    else:
        s = 1 if y2 > y1 else -1
        draw.polygon([(x2,y2),(x2-9,y2-15*s),(x2+9,y2-15*s)], fill="#24364B")

def centered(draw, box, lines, fnt):
    x1,y1,x2,y2 = box
    values = lines.split("\n")
    step = fnt.size + 8
    y = (y1+y2)/2 - step*(len(values)-1)/2
    for value in values:
        draw.text(((x1+x2)/2,y), value, font=fnt, fill="#10213B", anchor="mm")
        y += step

def make_er():
    im = Image.new("RGB", (1500, 900), "white"); d = ImageDraw.Draw(im)
    d.text((750,45), "ER-діаграма локальної бази даних", font=font(38,True), fill="#10213B", anchor="ma")
    d.rounded_rectangle((360,130,1140,760), 24, fill="#EAF4FF", outline="#315C9B", width=4)
    d.rectangle((360,130,1140,235), fill="#315C9B")
    d.text((750,182), "sensor_records", font=font(34,True), fill="white", anchor="mm")
    rows = [
        ("PK", "id", "INTEGER / Long", "AUTOINCREMENT"),
        ("IDX", "timestamp", "INTEGER / Long", "Unix ms"),
        ("", "sensorValue", "REAL / Float", "m/s²"),
        ("", "actionTriggered", "INTEGER / Boolean", "0 або 1"),
    ]
    y = 275
    for marker,name,typ,note in rows:
        d.text((420,y), marker, font=font(22,True), fill="#315C9B")
        d.text((540,y), name, font=font(25,True), fill="#10213B")
        d.text((790,y), typ, font=font(24), fill="#10213B")
        d.text((1010,y), note, font=font(22), fill="#37474F")
        d.line((400,y+48,1100,y+48), fill="#B7C7D9", width=2); y += 105
    im.save(ER)

def make_mvvm():
    im = Image.new("RGB", (1800, 1040), "white"); d = ImageDraw.Draw(im)
    d.text((900,35), "Фактична архітектура запису та аналітики", font=font(38,True), fill="#10213B", anchor="ma")
    boxes = [
        (50,170,330,285,"SensorManager"),(400,170,720,285,"SensorViewModel"),
        (790,150,1160,305,"LocalHistory\nRecorderViewModel"),(1230,150,1580,305,"SensorHistory\nRepository"),
        (1280,410,1580,525,"Room DAO"),(1280,660,1580,775,"SQLite"),
        (790,660,1160,775,"AnalyticsViewModel"),(400,660,720,775,"StateFlow"),(50,640,330,795,"AnalyticsScreen\nVico LineChart"),
    ]
    for x1,y1,x2,y2,label in boxes:
        d.rounded_rectangle((x1,y1,x2,y2),18,fill="#EAF4FF" if y1<400 else "#E9F8EF",outline="#315C9B",width=3)
        centered(d,(x1,y1,x2,y2),label,font(23,True))
    for start,end in [((330,227),(400,227)),((720,227),(790,227)),((1160,227),(1230,227)),((1405,305),(1405,410)),((1430,525),(1430,660)),((1280,717),(1160,717)),((790,717),(720,717)),((400,717),(330,717))]: arrow(d,start,end)
    d.text((900,900), "Successful actuator activation → pending flag → наступний успішний SensorRecord", font=font(24,True), fill="#7A3E00", anchor="ma")
    im.save(MVVM)

make_er(); make_mvvm()
if not SCREEN.exists(): raise FileNotFoundError(SCREEN)

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21); section.page_height = Cm(29.7)
section.top_margin = Cm(2); section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5); section.right_margin = Cm(1.5)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"; normal.font.size = Pt(14)
normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")

def set_font(run, name="Times New Roman", size=14, bold=False, italic=False, color=(0,0,0)):
    run.font.name=name; run.font.size=Pt(size); run.font.bold=bold; run.font.italic=italic; run.font.color.rgb=RGBColor(*color)
    rpr=run._element.get_or_add_rPr(); rpr.rFonts.set(qn("w:ascii"),name); rpr.rFonts.set(qn("w:hAnsi"),name); rpr.rFonts.set(qn("w:eastAsia"),name)
    return run

def p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, bold=False, italic=False, size=14, before=0, after=0, keep=False, spacing=1.5):
    para=doc.add_paragraph(); para.alignment=align; para.paragraph_format.line_spacing=spacing
    para.paragraph_format.space_before=Pt(before); para.paragraph_format.space_after=Pt(after)
    para.paragraph_format.first_line_indent=Cm(1.25) if indent else Cm(0); para.paragraph_format.keep_with_next=keep
    set_font(para.add_run(text),size=size,bold=bold,italic=italic); return para

def heading(text): return p(text,WD_ALIGN_PARAGRAPH.CENTER,False,True,size=14,before=8,after=4,keep=True)
def stage(text): return p(text,WD_ALIGN_PARAGRAPH.LEFT,False,True,size=14,before=7,after=3,keep=True)
def caption(text): return p(text,WD_ALIGN_PARAGRAPH.CENTER,False,italic=True,size=12,after=5,keep=True,spacing=1)
def page_break(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def bullet(text, number=None):
    para=p((f"{number}. " if number else "- ")+text,indent=False); para.paragraph_format.left_indent=Cm(.6); para.paragraph_format.first_line_indent=Cm(-.6); return para

def shade(cell, fill):
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); cell._tc.get_or_add_tcPr().append(shd)

def table(title, headers, rows, widths, size=10):
    caption(title); t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for cell,text,width in zip(t.rows[0].cells,headers,widths):
        cell.width=Cm(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(cell,"315C9B")
        cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(cell.paragraphs[0].add_run(text),size=size,bold=True,color=(255,255,255))
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for cell,text,width in zip(cells,row,widths):
            cell.width=Cm(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cell,"F2F6FA")
            para=cell.paragraphs[0]; para.paragraph_format.first_line_indent=Cm(0); para.paragraph_format.line_spacing=1
            set_font(para.add_run(str(text)),size=size)
    p("",indent=False,size=4,spacing=1); return t

def code(title,text,size=8.2):
    caption(title); t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; t.autofit=False; t.columns[0].width=Cm(16.2)
    cell=t.cell(0,0); shade(cell,"F3F5F7"); para=cell.paragraphs[0]; para.paragraph_format.first_line_indent=Cm(0); para.paragraph_format.line_spacing=1
    for i,line in enumerate(text.strip().splitlines()):
        if i: para.add_run().add_break()
        set_font(para.add_run(line),"Consolas",size)
    p("",indent=False,size=4,spacing=1)

def image(path,width,cap,descr):
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(path),width=Cm(width))
    for pr in para._p.xpath(".//wp:docPr"): pr.set("title",cap); pr.set("descr",descr)
    p(cap,WD_ALIGN_PARAGRAPH.CENTER,False,italic=True,size=12,after=5,keep=False,spacing=1)

# Title page
for text,size,bold in [
    ("НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ «ЛЬВІВСЬКА ПОЛІТЕХНІКА»",14,True),
    ("Кафедра автоматизованих систем управління",14,False),
]: p(text,WD_ALIGN_PARAGRAPH.CENTER,False,bold=bold,size=size,spacing=1)
if LOGO.exists(): image(LOGO,4.0,"","Логотип Національного університету Львівська політехніка")
p("ЗВІТ",WD_ALIGN_PARAGRAPH.CENTER,False,True,size=18,before=18,spacing=1)
p("до лабораторної роботи № 5",WD_ALIGN_PARAGRAPH.CENTER,False,True,size=16,spacing=1)
p("Локальне зберігання даних у БД Room та візуалізація динамічних процесів",WD_ALIGN_PARAGRAPH.CENTER,False,True,italic=True,size=15,before=8,spacing=1.15)
p("з дисципліни «Програмування мобільних смарт-систем»",WD_ALIGN_PARAGRAPH.CENTER,False,size=14,before=8,spacing=1)
p("",indent=False,size=14,before=35)
p("Виконав: студент групи ОІ-43\nГончаров Вадим",WD_ALIGN_PARAGRAPH.RIGHT,False,size=14,spacing=1.2)
p("Перевірив: Зварич Р.О.",WD_ALIGN_PARAGRAPH.RIGHT,False,size=14,before=10,spacing=1.2)
p("Львів 2026",WD_ALIGN_PARAGRAPH.CENTER,False,size=14,before=55,spacing=1)
page_break()

heading("Мета роботи")
p("Опанувати локальне зберігання часових рядів у Room, організувати асинхронний запис обробленої інтенсивності руху в архітектурі MVVM та побудувати реактивний лінійний графік за даними SQLite без залежності від мережі.")
heading("Завдання на лабораторну роботу")
for i,text in enumerate([
    "Створити Entity SensorRecord, DAO та singleton Room Database з експортом схеми.",
    "Додати Repository і recorder, який записує валідне значення приблизно раз на секунду.",
    "Зафіксувати лише нові успішні активації актуаторів у полі actionTriggered.",
    "Створити AnalyticsViewModel і окремий Compose-екран з графіком Vico.",
    "Обмежити графік 50 найновішими точками за останню годину.",
    "Перевірити unit, instrumented, lint, build та runtime-сценарії.",
],1): bullet(text,i)

heading("Технічний стек")
table("Таблиця 1. Використані компоненти",["Компонент","Версія або роль"],[
    ("Kotlin / Coroutines / Flow","2.3.21 / 1.10.2 / асинхронний та реактивний конвеєр"),
    ("Room","2.8.5: Entity, DAO, Database, KTX"),
    ("KSP","2.3.12: генерація Room implementation і JSON schema"),
    ("Jetpack Compose / Material 3","Екран Dashboard, Analytics і навігація"),
    ("Vico","3.3.1: Compose LineChart"),
    ("JUnit / AndroidX Test","99 unit tests і 3 instrumented tests"),
], [5.3,10.9],10.2)

heading("Хід виконання роботи")
stage("Етап 1. Схема Room")
p("База sleep_tracker.db має версію 1. Єдина таблиця sensor_records містить обов’язкові поля методички. Поле timestamp зберігає Unix time у мілісекундах, а не монотонний SensorEvent.timestamp. Для часових запитів Room створює індекс index_sensor_records_timestamp.")
image(ER,14.8,"Рис. 1. ER-діаграма таблиці sensor_records","Поля, типи та індекс таблиці локальної історії")
code("Лістинг 1. Entity SensorRecord",'''@Entity(tableName = "sensor_records", indices = [Index(value = ["timestamp"])])
data class SensorRecord(
    @PrimaryKey(autoGenerate = true) val id: Long = 0,
    val timestamp: Long,
    val sensorValue: Float,
    val actionTriggered: Boolean,
)''')

stage("Етап 2. DAO та порядок часових даних")
p("DAO спочатку вибирає найновіші рядки через ORDER BY timestamp DESC, id DESC і LIMIT. Зовнішній SELECT повертає відібрані рядки за зростанням часу, тому графік не розвертає часовий ряд. Repository обмежує limit діапазоном 1..500 і перевіряє timestamp, NaN, Infinity та від’ємні значення.")
code("Лістинг 2. Запит останніх N вимірювань",'''SELECT * FROM (
    SELECT * FROM sensor_records
    WHERE timestamp >= :sinceInclusive
    ORDER BY timestamp DESC, id DESC
    LIMIT :limit
) ORDER BY timestamp ASC, id ASC''')
table("Таблиця 2. Операції SensorRecordDao",["Операція","Призначення"],[
    ("insert", "Асинхронне додавання SensorRecord."),
    ("observeRecent", "Flow останніх N рядків у хронологічному порядку."),
    ("recordsBetween", "Вибірка заданого часового інтервалу."),
    ("observeCount", "Реактивна кількість вимірювань."),
    ("observeActionCount", "Реактивна кількість actionTriggered = true."),
    ("deleteOlderThan", "Безпечна retention-операція за timestamp."),
], [5.2,11.0],10.2)

stage("Етап 3. MVVM та Repository")
image(MVVM,16.0,"Рис. 2. MVVM-конвеєр локальної історії","SensorViewModel, recorder, Repository, Room, AnalyticsViewModel і Compose")
p("Room не викликається з Composable. SensorViewModel залишається єдиним джерелом MovementAnalysis. LocalHistoryRecorderViewModel спостерігає його StateFlow, тому перехід між Dashboard і Analytics не створює другого SensorManager listener та не зупиняє запис. Singleton Database отримує applicationContext і не використовує allowMainThreadQueries або destructive migration.")

stage("Етап 4. Контроль частоти запису")
p("Recorder має один coroutine loop у viewModelScope. Раз на 1000 ms він бере останній валідний Active state і пропускає повтор, якщо timestampNanos сенсорного sample не змінився. Для Room формується окремий System.currentTimeMillis(). Запис виконується Repository на Dispatchers.IO.")
code("Лістинг 3. Запис нового вимірювання",'''if (lastRecordedSensorTimestampNanos == sensorTimestamp) return
writeMutex.withLock {
    val actionForRecord = pendingActivation
    repository.addMeasurement(clock(), value, actionForRecord)
    lastRecordedSensorTimestampNanos = sensorTimestamp
    if (actionForRecord) pendingActivation = false
}''')

stage("Етап 5. Семантика подій актуаторів")
p("ActuatorViewModel надсилає ActuatorActivation через необмежений Channel лише після ActuatorOperationStatus.SUCCESS. Подіями є VIBRATION_PULSE, VIBRATION_ENABLED і FLASHLIGHT_ENABLED. Вимкнення, UNAVAILABLE, ERROR та повторне застосування вже активного persistent state не вважаються новою активацією. Recorder тримає pendingActivation до успішного insert, тому помилка Room не втрачає подію.")

stage("Етап 6. Екран Analytics")
p("AnalyticsViewModel кожну хвилину оновлює межу rolling one-hour window і через flatMapLatest підписується на три Room Flow. Графік отримує не більше 50 найновіших точок. X використовує фактичні нерівномірні timestamps як відносні секунди від першої точки, що усуває втрату точності Double; formatter відновлює локальний час. Y показує smoothed movement intensity у m/s².")
image(SCREEN,5.2,"Рис. 3. Екран Analytics з реальними записами Room","Runtime screenshot з емулятора Pixel 10 Pro, графік рухової активності та лічильники")
p("На фактичному скриншоті відображено 394 вимірювання і 17 успішних активацій після instrumented test run та повторного запуску. Видимі піки створено командами virtual accelerometer. Дані залишалися доступними при BLE disconnected; Firebase не бере участі у локальному записі.")

stage("Етап 7. Перевірка")
table("Таблиця 3. Фактично виконані перевірки",["Перевірка","Результат"],[
    ("testDebugUnitTest", "99 passed, 0 failed"),
    ("connectedDebugAndroidTest", "3 passed, 0 failed; Pixel_10_Pro, Android 16"),
    ("assembleDebug", "BUILD SUCCESSFUL"),
    ("lintDebug", "BUILD SUCCESSFUL; 0 errors"),
    ("compileDebugAndroidTestSources", "BUILD SUCCESSFUL"),
    ("Runtime Analytics", "Room створено; 394 записи, 17 activations; Vico оновлюється"),
    ("Process restart", "Історія збережена після force-stop/start без clear data"),
], [6.2,10.0],10.2)
p("Під час першого runtime запуску Vico завершувався з IllegalArgumentException через надто точні абсолютні X-values. Перехід на відносні секунди усунув crash. Окремо додано Material 3 Vico theme provider і background Analytics для читабельного dark-theme контрасту.")

stage("Статус доказів")
p("Runtime screenshot Analytics, generated Room schema, unit/instrumented results і факт створення sleep_tracker.db підтверджено. Скріншот Android Studio Database Inspector не створено автоматично, оскільки він має походити саме з інтерактивного App Inspection UI. До остаточного звіту потрібно додати один такий кадр з відкритою таблицею sensor_records і видимими рядками actionTriggered.")

heading("Висновки")
p("Room практичніший за пряме керування SQLite, оскільки перевіряє SQL під час компіляції, генерує типобезпечний DAO, інтегрується з suspend і Flow та централізує версію схеми. Це зменшує кількість ручного Cursor-коду і помилок перетворення типів.")
p("Локальна база підвищує автономність Sleep Tracker Simulator: рухова історія записується при відсутньому BLE та незалежно від Firebase або Інтернету, а SQLite зберігає її після завершення процесу. Room Flow автоматично перевидає дані після insert, тому AnalyticsViewModel і Vico оновлюють графік без polling та без запитів із Composable.")
p("Історію актуаторів інтегровано як одноразові типізовані події успішної активації. Mutex зв’язує pending event з наступним успішним SensorRecord, а retry-семантика не очищає pending flag після невдалого insert. Фактичні перевірки завершилися 99 успішними unit tests, 3 успішними instrumented tests, успішними build і lint та runtime-відображенням реальних записів. Документ залишається робочою версією лише через відсутній скриншот Database Inspector.")

doc.core_properties.title = "Лабораторна робота 5 Room та аналітика"
doc.core_properties.subject = "Програмування мобільних смарт-систем"
doc.core_properties.author = "Гончаров Вадим"
doc.save(OUT)
print(OUT)
