from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


PATH = Path(r"C:\Users\Eclipse\Desktop\7 семестр\Програмування мобільних смарт-систем\ЛР-1_МПСС_Гончаров_Вадим_ОІ-43_final.docx")
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def main():
    with zipfile.ZipFile(PATH) as archive:
        bad = archive.testzip()
        if bad:
            raise RuntimeError(f"Corrupt ZIP member: {bad}")
        xml = archive.read("word/document.xml")
        root = etree.fromstring(xml)
        equations = root.xpath("//m:oMath", namespaces=NS)
        math_paras = root.xpath("//m:oMathPara", namespaces=NS)
        drawings = root.xpath("//w:drawing", namespaces=NS)
        document_text = "".join(root.xpath("//w:t/text()", namespaces=NS))

    doc = Document(PATH)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    all_text = "\n".join(paragraphs + [cell.text for t in doc.tables for row in t.rows for cell in row.cells])

    headings = [
        "Мета роботи",
        "Концепція наскрізного завдання",
        "Завдання на лабораторну роботу",
        "Технічний стек",
        "Індивідуальне завдання",
        "Обґрунтування порогових значень",
        "Хід роботи",
        "Висновок",
    ]
    positions = {heading: paragraphs.index(heading) for heading in headings}
    if list(positions.values()) != sorted(positions.values()):
        raise RuntimeError("Required sections are out of order")

    caption_numbers = [int(n) for n in re.findall(r"Рис\.\s*(\d+)\.", all_text)]
    listing_numbers = [int(n) for n in re.findall(r"Лістинг\s*(\d+)\.", all_text)]

    checks = {
        "zip_integrity": bad is None,
        "native_oMath_count": len(equations),
        "oMathPara_count": len(math_paras),
        "equation_markers_remaining": document_text.count("WORD_EQUATION"),
        "drawings_total": len(drawings),
        "inline_shapes_python_docx": len(doc.inline_shapes),
        "code_tables": len(doc.tables),
        "figure_captions": caption_numbers,
        "listing_captions": listing_numbers,
        "em_dash_count": all_text.count("—"),
        "en_dash_count": all_text.count("–"),
        "variant_4_count": all_text.count("Варіант №4") + all_text.count("варіант №4"),
        "other_variant_count": len(re.findall(r"[Вв]аріант №(?!4\b)\d+", all_text)),
        "forbidden_terms": {term: all_text.count(term) for term in ["Firebase", "Room", "TFLite", "BLE"]},
        "section_positions": positions,
    }
    print(checks)

    assert len(equations) == 2
    assert document_text.count("WORD_EQUATION") == 0
    assert len(doc.inline_shapes) == 6  # title logo plus five report screenshots
    assert len(doc.tables) == 3
    assert caption_numbers == [1, 2, 3, 4, 5]
    assert listing_numbers == [1, 2, 3]
    assert all_text.count("—") == 0 and all_text.count("–") == 0
    assert checks["other_variant_count"] == 0
    assert all(value == 0 for value in checks["forbidden_terms"].values())


if __name__ == "__main__":
    main()
