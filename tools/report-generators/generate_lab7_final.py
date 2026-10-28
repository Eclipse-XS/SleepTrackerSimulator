from pathlib import Path
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from generate_labs_6_7_reports import reset_reference, title, h, para, code, table

ROOT = Path(__file__).resolve().parents[2]

def add_picture(doc, path, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(8.8))
    c = doc.add_paragraph(caption); c.alignment = WD_ALIGN_PARAGRAPH.CENTER

def build():
    d = reset_reference()
    title(d, 7, "Мініпроєкт: складання та тестування цілісної смарт-системи")
    h(d, "1 Тема, мета та індивідуальний домен")
    para(d, "Тема: інтеграція компонентів мобільної смарт-системи та фінальна перевірка. Мета — об’єднати реальний Android SensorManager, локальний TensorFlow Lite inference, Room, Firebase, WorkManager, локальні актуатори та Mock BLE. Індивідуальний домен — Sleep Tracker Simulator. Система має навчальне призначення і не є медичним засобом.")
    h(d, "2 Функціональні та нефункціональні вимоги")
    table(d,["Група","Вимога","Реалізація"],[
      ("Input","XYZ та movement intensity","AndroidAccelerometerDataSource"),("AI","REST/RESTLESS/ACTIVE","TFLite [1,20,3] → [1,3]"),("Storage","History, rules, executions, outbox","Room schema v4"),("Cloud","Durable Firebase delivery","Outbox + WorkManager"),("IoT","HR та alert command","VirtualSleepBand Mock BLE"),("Output","Vibration, flashlight, visible alert","AndroidSmartActuator / AlertDialog"),("Quality","Lifecycle, retry, validation","ViewModels, Flow, tests")],[3.2,6.4,6.4])
    para(d,"Нефункціональні вимоги: відсутність блокувального I/O на main thread, збереження правил після process death, явні error/pending states, owner-scoped Firebase rules, контроль життєвого циклу sensor listener та відтворювані тести.")
    h(d,"3 Архітектура і потік даних")
    code(d,"SensorManager -> SensorViewModel -> Compose Dashboard\n        |-> Room sensor history -> Last five records\n        |-> 20x3 buffer -> TFLite -> Smart Rule Engine\nSmart Rule Engine -> Room execution\n                  |-> AndroidSmartActuator\n                  |-> MockBleConnector ALERT\nMonitoring snapshot -> Room outbox -> WorkManager -> Firebase")
    para(d,"Гілки Room, cloud, actuator і BLE асинхронні. Вони не утворюють одну розподілену транзакцію. Coordinator фіксує фактичний результат кожної rule action окремо.")
    h(d,"4 Smart Scenarios")
    para(d,"Користувач створює, редагує, вмикає, вимикає і видаляє правила. Sources охоплюють X/Y/Z, movement intensity та ML class. Numeric comparisons підтримують >, >=, <, <=. ML source допускає class equality. Invalid combinations відхиляються до persistence.")
    table(d,["Поле","Значення прикладу"],[('name','X threshold alert'),('source','ACCELEROMETER_X, m/s²'),('operator','>'),('threshold','10.0'),('action','VIBRATION_PULSE'),('requiredConsecutiveMatches','3'),('cooldownMillis','5000')],[7,9])
    para(d,"Pure SmartRuleEngine не виконує Android side effects. Він веде consecutive count, armed state і monotonic cooldown. Тривале true condition дає один trigger, доки правило не rearm після false condition.")
    add_picture(d, ROOT/'docs/evidence/lab7/smart-scenario-saved.png', "Рисунок 1 — Збережене користувацьке правило на AVD")
    h(d,"5 Машинне навчання")
    para(d,"ActivityClassifier завантажує реальний activity_classifier.tflite. Вхідний tensor має форму [1,20,3], вихідний — [1,3]. Класи REST, RESTLESS і ACTIVE описують синтетичні патерни руху. Незалежний synthetic test set містить 900 вікон і дав accuracy 97,22%. Це не клінічна точність.")
    h(d,"6 Room і міграції")
    para(d,"Schema v4 додає smart_rules та smart_rule_executions з foreign key і cascade delete. DAO повертає Flow, mutation виконується на Dispatchers.IO. Instrumentation tests перевіряють CRUD, execution history, cascade та ланцюги 2→4 і 1→2→3→4 без втрати sensor history.")
    h(d,"7 Firebase і offline recovery")
    para(d,"CloudSyncViewModel спочатку додає stable operation до Room outbox. Firebase acknowledgement переводить її у delivered. За відсутності мережі operation лишається pending; WorkManager повторює delivery. Capacity telemetry queue — 10000, terminal END_SESSION має окрему безпечну семантику. Smart-rule execution поки не має власного cloud payload, тому це заявлено як обмеження.")
    h(d,"8 Mock BLE та актуатори")
    para(d,"VirtualSleepBand моделює scan, connection, reads, writes і HR notifications. DeviceCommand.ALERT обробляється лише у connected state та повертає ALERT_ACKNOWLEDGED. Vibration pulse і flashlight проходять AndroidSmartActuator; flashlight потребує permission і наявного hardware. Visible action показує Material AlertDialog.")
    h(d,"9 Dashboard")
    para(d,"Dashboard показує BLE connection, heart rate, XYZ, movement intensity, TFLite class/confidence, Firebase/outbox state, quick actuator controls і п’ять найновіших SensorRecord, прочитаних з Room Flow.")
    h(d,"10 Фільтрація хибних спрацювань")
    para(d,"Фільтрація складається з трьох рівнів: рухова метрика домену, windowed ML inference та rule-level consecutive matches/rising edge/cooldown. Exact-boundary та sustained-condition cases покриті unit tests. Runtime rearm state обнуляється після process restart.")
    h(d,"11 Інтеграційне тестування")
    table(d,["ID","Перевірка","Фактичний результат"],[
      ("L7-01","Rule: 3 matches, edge, cooldown","PASS, JVM"),("L7-02","Rule invalid combinations/boundaries","PASS, JVM"),("L7-03","Mock BLE ALERT acknowledgement","PASS, JVM"),("L7-04","Room rule CRUD/cascade","PASS, AVD instrumentation"),("L7-05","Room migrations through v4","PASS, AVD instrumentation"),("L7-06","Real TFLite load/inference","PASS, AVD instrumentation"),("L7-07","Rule survives app restart","PASS, AVD screenshot/video"),("L7-08","Full desktop E2E with Firebase console","PARTIAL, not recorded")],[2,7,7])
    h(d,"12 Профілювання")
    table(d,["Точка","TOTAL PSS","TOTAL RSS"],[('Start','115111 KB','248764 KB'),('30 s','130093 KB','275868 KB'),('60 s','129377 KB','275564 KB'),('10 restarts','92940 KB','210920 KB')],[5,5.5,5.5])
    para(d,"У 60-секундному debug AVD run swap дорівнював 0. PSS між 30 і 60 секундами зменшився на 716 KB, тобто монотонного росту в цьому інтервалі не виявлено. Gfxinfo: 1028 кадрів, 92,02% janky, p50 20 ms, p90 65 ms, p95 101 ms, p99 250 ms. Це незадовільний UX результат і причина production NO-GO до release profiling на фізичному пристрої.")
    h(d,"13 Виняткові ситуації")
    table(d,["Ситуація","Поведінка"],[('Network/Firebase failure','Pending outbox, retry, visible error'),('Sensor unavailable','Unavailable state, без вигаданих samples'),('BLE disconnected','ALERT error, без false delivery'),('Model init failure','ML error state, UI не падає'),('Outbox full','Backpressure; pending telemetry не eviction'),('Local DB failure','Repository error, action не видається за persisted'),('Activity recreation','ViewModel/Room state відновлюється'),('Flashlight unavailable','Unsupported/permission result')],[5,11])
    h(d,"14 Відеодоказ")
    para(d,"Файл docs/evidence/lab7/lab7-smart-scenario.mp4 створено справжнім adb shell screenrecord. Він показує UI створення та persistence правила. Він не охоплює Extended Controls і Firebase Console, тому не вважається повним доказом TC-L7-28.")
    h(d,"15 Висновки")
    para(d,"Lab 7 має повну навчальну реалізацію persistent Smart Scenarios, Room v4, BLE alert command, Dashboard recent records та real TFLite pipeline. Автоматизовані тести підтверджують core contracts. Статус захисту — READY WITH LIMITATIONS через часткове відео, відсутність окремої cloud-події smart execution і високий jank debug AVD. Production verdict — NO-GO до фізичного profiling, реального BLE, security/privacy hardening і валідації ML на реальних даних.")
    out=ROOT/'deliverables/ЛР-7_ПМСС_Гончаров_Вадим_ОІ-43.docx'; d.save(out); print(out)

if __name__ == '__main__': build()
